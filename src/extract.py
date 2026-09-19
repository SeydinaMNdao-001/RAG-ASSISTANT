"""
extract.py — Extraction de texte à partir de différents formats de documents.

- PDF avec calque texte -> extraction directe
- PDF scanné (image) -> bascule automatique vers l'OCR, page par page
- DOCX (contrats Word) -> paragraphes + tableaux
- Images seules (PNG/JPG) -> OCR direct
- TXT / MD -> lecture directe

Chaque fonction retourne une liste de (texte, numéro_de_page_ou_None).
"""

from __future__ import annotations

import io
from pathlib import Path

import fitz  # PyMuPDF
import pytesseract
from docx import Document as DocxDocument
from PIL import Image
from pypdf import PdfReader

from src.config import OCR_ENABLED, OCR_LANGUAGES, OCR_MIN_TEXT_LENGTH, OCR_RENDER_DPI


def _ocr_pdf_page(pdf_path: Path, page_index: int) -> str:
    """Rend une page de PDF en image puis l'envoie à Tesseract OCR."""
    doc = fitz.open(str(pdf_path))
    try:
        page = doc[page_index]
        zoom = OCR_RENDER_DPI / 72  # 72 dpi = résolution native d'un PDF
        pix = page.get_pixmap(matrix=fitz.Matrix(zoom, zoom))
        image = Image.open(io.BytesIO(pix.tobytes("png")))
        return pytesseract.image_to_string(image, lang=OCR_LANGUAGES)
    finally:
        doc.close()


def extract_pdf(path: Path) -> list[tuple[str, int]]:
    """Extrait le texte d'un PDF, page par page, avec bascule OCR si besoin."""
    reader = PdfReader(str(path))
    pages: list[tuple[str, int]] = []

    for i, page in enumerate(reader.pages):
        text = (page.extract_text() or "").strip()

        if len(text) < OCR_MIN_TEXT_LENGTH and OCR_ENABLED:
            ocr_text = _ocr_pdf_page(path, i).strip()
            if len(ocr_text) > len(text):
                text = ocr_text

        if text:
            pages.append((text, i + 1))

    return pages


def extract_docx(path: Path) -> list[tuple[str, None]]:
    """Extrait le texte d'un document Word (.docx), paragraphes + tableaux."""
    doc = DocxDocument(str(path))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]

    for table in doc.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells if c.text.strip()]
            if cells:
                parts.append(" | ".join(cells))

    text = "\n\n".join(parts)
    return [(text, None)] if text.strip() else []


def extract_image(path: Path) -> list[tuple[str, None]]:
    """OCR direct sur une image (photo de document, scan exporté en JPG/PNG...)."""
    if not OCR_ENABLED:
        return []
    image = Image.open(path)
    text = pytesseract.image_to_string(image, lang=OCR_LANGUAGES)
    return [(text, None)] if text.strip() else []


def extract_plain_text(path: Path) -> list[tuple[str, None]]:
    """Fichiers texte brut (.txt, .md)."""
    text = path.read_text(encoding="utf-8", errors="ignore")
    return [(text, None)] if text.strip() else []


EXTRACTORS = {
    ".pdf": extract_pdf,
    ".docx": extract_docx,
    ".txt": extract_plain_text,
    ".md": extract_plain_text,
    ".png": extract_image,
    ".jpg": extract_image,
    ".jpeg": extract_image,
}


def extract_document(path: Path) -> list[tuple[str, int | None]]:
    """Point d'entrée unique : dispatch vers le bon extracteur selon l'extension."""
    extractor = EXTRACTORS.get(path.suffix.lower())
    if extractor is None:
        raise ValueError(f"Format non supporté : {path.suffix}")
    return extractor(path)